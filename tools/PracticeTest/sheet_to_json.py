# coding: utf-8
"""연습전투 테스트 시트(세나테스트용.xlsx) → JSON 변환.

사용법: python3 tools/PracticeTest/sheet_to_json.py 세나테스트용.xlsx tools/PracticeTest/team.json

시트 1개 = 영웅 1명 (시트 이름 무관, A3에 유형이 있는 시트만). 셀 위치는 템플릿 고정:
  3행: 유형/등급/이름/초월/잠재(공방생)/스킬강화(평타,1스킬,2스킬,패시브)
  A6: 세트명(4부위 공통), C~N 6~12행: 무기1/무기2/방어구1/방어구2 (메인·부옵1~4)
  C15~D19: 장신구(성급보너스 3줄 + 메인 + 부옵), E15~G19: 전용장비(공격력 + 조율1~4)
수치 규칙: 1 미만 = 비율(0.24 → 24%), 1 이상 = 깡수치(50 → 공격력 50). 부옵은 티어 수로 환산·검증.
"""
import json
import sys

import openpyxl

# 시트 표기 → 계산기 스탯 키 (비율일 때 / 깡수치일 때)
STAT_NAMES = {
    "모든공격력": ("공격력%", "공격력"), "공격력": ("공격력%", "공격력"),
    "방어력": ("방어력%", "방어력"), "생명력": ("생명력%", "생명력"),
    "치명타확률": ("치명타확률%", None), "치명타피해": ("치명타피해%", None),
    "약점공격확률": ("약점공격확률%", None), "막기확률": ("막기확률%", None),
    "효과적중": ("효과적중%", None), "효과저항": ("효과저항%", None),
    "받는피해감소": ("받피감%", None), "속공": (None, "속공"),
}
SUB_TIER = {"공격력%": 5, "공격력": 50, "치명타확률%": 4, "치명타피해%": 6, "약점공격확률%": 5, "속공": 4,
            "막기확률%": 4, "효과적중%": 5, "효과저항%": 5, "방어력%": 5, "방어력": 30, "생명력%": 5, "생명력": 180}
ACC_NAMES = {"보스피해증가": "보피증%", "1~3인기피해증가": "1-3인기%", "4~5인기피해증가": "4-5인기%",
             "피해증가": "피증%", "치명타확률": "치명타확률%", "약점공격확률": "약점공격확률%", "막기확률": "막기%",
             "방어력": "방어력%", "생명력": "생명력%", "효과적중": "효과적중%", "효과저항": "효과저항%"}
ACC_GRADE = {0.05: 4, 0.07: 5, 0.1: 6}
TUNING = {"모든공격력": "모든공격력", "방어력": "방어력", "생명력": "생명력", "효과적중": "효과적중",
          "효과저항": "효과저항", "피해증폭": "피해증폭", "파쇄": "파쇄", "탄성": "탄성"}
# 조율 수치 → 등급 (ExclusiveWeapon.GetTuningValue 표 역산)
TUNING_GRADE = {
    "모든공격력": {5: "고급", 7: "희귀", 12: "전설"}, "방어력": {5: "고급", 7: "희귀", 12: "전설"},
    "생명력": {5: "고급", 7: "희귀", 12: "전설"}, "효과적중": {4: "고급", 6: "희귀", 10: "전설"},
    "효과저항": {4: "고급", 6: "희귀", 10: "전설"}, "피해증폭": {1.6: "고급", 2.4: "희귀", 4: "전설"},
    "파쇄": {4.8: "고급", 7.2: "희귀", 12: "전설"}, "탄성": {6: "고급", 9: "희귀", 15: "전설"},
}


TYPO = {"방여력": "방어력", "생명렵": "생명력"}
# 특수 장신구(반지): 효과 문구 키워드 → 반지 이름. 반지는 메인옵 자리를 대신한다(그 다음 줄 = 부옵).
RING_KEYWORDS = [("권능", "권능의 반지"), ("부활", "부활의 반지"), ("불사", "불사의 반지"), ("화상", "샐러맨더의 반지"),
                 ("출혈", "가시 반지"), ("중독", "독사 반지"), ("기절", "기회의 반지"), ("침묵", "저주의 반지")]


def stat_key(name, value, warn):
    name = str(name).replace(" ", "")
    name = TYPO.get(name, name)
    pair = STAT_NAMES.get(name)
    if pair is None:
        warn.append(f"알 수 없는 스탯 이름: {name}")
        return None, None
    is_rate = value < 1
    key = pair[0] if is_rate else pair[1]
    if key is None:
        warn.append(f"{name} {value}: {'비율' if is_rate else '깡수치'} 표기가 이 스탯엔 없음")
        return None, None
    return key, (round(value * 100, 4) if is_rate else value)


def read_piece(ws, col, slot, set_name, warn):
    """col = 'C'(무기1) / 'F'(무기2) / 'I'(방어구1) / 'L'(방어구2). 이름칸=col+1, 값칸=col+2."""
    c0 = openpyxl.utils.column_index_from_string(col)
    name_col, val_col = c0 + 1, c0 + 2
    main_name, main_val = ws.cell(8, name_col).value, ws.cell(8, val_col).value
    mk, mv = stat_key(main_name, main_val, warn)
    subs = []
    for r in range(9, 13):
        n, v = ws.cell(r, name_col).value, ws.cell(r, val_col).value
        if n is None:
            continue
        k, val = stat_key(n, v, warn)
        if k is None:
            continue
        per = SUB_TIER.get(k)
        tier = val / per if per else 0
        if per is None or abs(tier - round(tier)) > 1e-6:
            warn.append(f"{slot} 부옵 {n} {v}: 티어 환산 불가({val}/{per})")
        subs.append({"stat": k, "value": val, "tier": int(round(tier))})
    return {"slot": "무기" if slot.startswith("무기") else "방어구", "name": slot, "set": set_name,
            "main": mk, "mainValue": mv, "subs": subs}


def read_hero(ws):
    warn = []
    g = lambda a: ws[a].value
    hero = {
        "sheet": ws.title, "type": g("A3"), "grade": g("B3"), "name": g("C3"), "transcend": int(g("D3") or 0),
        "potential": {"atk": int(g("E3") or 0), "def": int(g("F3") or 0), "hp": int(g("G3") or 0)},
        "enhanced": {"normal": g("H3") == 1, "skill1": g("I3") == 1, "skill2": g("J3") == 1, "passive": g("K3") == 1},
    }
    set_name = g("A6")
    hero["gear"] = [read_piece(ws, c, s, set_name, warn)
                    for c, s in (("C", "무기1"), ("F", "무기2"), ("I", "방어구1"), ("L", "방어구2"))]
    # 장신구: C15~C17 성급 보너스(공/방/생 동일값) → 성급. C18~ : [메인 또는 반지 효과], [부옵]
    bonus = g("D15")
    acc = {"grade": ACC_GRADE.get(round(bonus or 0, 2), 0) if bonus else 0}
    if bonus and acc["grade"] == 0:
        warn.append(f"장신구 성급 보너스 {bonus} → 성급 판별 불가")
    if not bonus:
        warn.append("장신구 없음(빈칸)")
    lines = [(ws.cell(r, 3).value, ws.cell(r, 4).value) for r in range(18, 21) if ws.cell(r, 3).value is not None]
    first = True
    for text, val in lines:
        t = str(text).replace(" ", "")
        ring = next((rn for kw, rn in RING_KEYWORDS if kw in t), None)
        if ring and first:
            acc["ring"] = ring
            acc["ringValue"] = val
            first = False
            continue
        key = ACC_NAMES.get(t)
        if key is None:
            warn.append(f"장신구 줄 미인식: {text}")
            continue
        slot = "main" if first and "ring" not in acc else "sub"
        acc[slot] = key
        acc[slot + "Value"] = round(val * 100, 4)
        first = False
    hero["accessory"] = acc
    # 전용장비: E열 라벨 / F열(이름 또는 수치) / G열(조율 수치). 라벨 '조율N' = 조율, 그 외 = 깡스탯(공격력/생명력/방어력)
    ex = {"atk": 0, "hp": 0, "def": 0, "tuning": []}
    has_ex = False
    for r in range(15, 22):
        label = ws.cell(r, 5).value
        if label is None:
            continue
        label = str(label).replace(" ", "")
        has_ex = True
        if label.startswith("조율"):
            n, v = ws.cell(r, 6).value, ws.cell(r, 7).value
            if n is None or v is None:
                warn.append(f"{label} 빈칸")
                continue
            opt = TUNING.get(str(n).replace(" ", ""))
            pct = round(v * 100, 4)
            grade = TUNING_GRADE.get(opt, {}).get(pct if pct != int(pct) else int(pct))
            if opt is None or grade is None:
                warn.append(f"조율 {n} {v}: 옵션/등급 판별 불가")
            ex["tuning"].append({"option": opt, "value": pct, "grade": grade})
        else:
            key = {"공격력": "atk", "생명력": "hp", "방어력": "def"}.get(label)
            if key is None:
                warn.append(f"전용장비 줄 미인식: {label}")
            else:
                ex[key] = ws.cell(r, 6).value or 0
    if not has_ex:
        warn.append("전용장비 없음(빈칸)")
        ex = None
    hero["exclusive"] = ex
    hero["warnings"] = warn
    return hero


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)
    hero_types = {"공격형", "마법형", "방어형", "지원형", "만능형"}
    heroes = []
    for ws in wb.worksheets:
        if ws["A3"].value not in hero_types:
            continue
        h = read_hero(ws)
        # 시트명 "진형번호N" → 자리 번호 N
        digits = "".join(ch for ch in ws.title if ch.isdigit())
        h["slot"] = int(digits) if digits else None
        heroes.append(h)
    heroes.sort(key=lambda h: (h["slot"] is None, h["slot"] or 0))
    pet = None
    if "펫" in wb.sheetnames:
        ws = wb["펫"]
        pots = [ws.cell(3, c).value for c in range(4, 8)]
        pet = {"grade": ws["A3"].value, "name": ws["B3"].value, "enhance": ws["C3"].value,
               "potentialAtkRate": round(sum(v for v in pots if v) * 100, 4)}
        print(f"펫: {pet['name']} 스킬강화{pet['enhance']} 잠재 모공 합 {pet['potentialAtkRate']}%")
    json.dump({"heroes": heroes, "pet": pet}, open(dst, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    for h in heroes:
        print(f"[자리{h['slot']}] {h['sheet']}: {h['name']} 초월{h['transcend']} 세트 {h['gear'][0]['set']}"
              + (f"  경고 {len(h['warnings'])}건: " + " / ".join(h["warnings"]) if h["warnings"] else "  (경고 없음)"))


if __name__ == "__main__":
    main()
